#!/usr/bin/env python3
"""Word document generation helper using python-docx."""
import argparse
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocxBuilder:
    """Simple DOCX document builder."""

    def __init__(self, filename, title="Document"):
        self.filename = filename
        self.title = title
        self.document = Document()
        self.document.add_heading(title, level=0)

    def add_heading(self, text, level=1):
        self.document.add_heading(text, level=level)

    def add_paragraph(self, text, bold=False, italic=False):
        p = self.document.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic

    def add_table(self, data):
        """data: list of lists. First row is treated as header."""
        if not data:
            return
        table = self.document.add_table(rows=len(data), cols=len(data[0]))
        table.style = "Table Grid"
        for i, row in enumerate(data):
            for j, cell_text in enumerate(row):
                table.rows[i].cells[j].text = str(cell_text)
        # Bold header row
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    def add_bullet_list(self, items):
        for item in items:
            self.document.add_paragraph(item, style="List Bullet")

    def add_numbered_list(self, items):
        for item in items:
            self.document.add_paragraph(item, style="List Number")

    def add_page_break(self):
        self.document.add_page_break()

    def save(self):
        self.document.save(self.filename)
        return self.filename


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate a Word document")
    parser.add_argument("--output", default="output.docx", help="Output DOCX path")
    parser.add_argument("--title", default="Document", help="Document title")
    args = parser.parse_args()

    doc = DocxBuilder(args.output, title=args.title)
    doc.add_paragraph("This is a generated document. Add content programmatically using DocxBuilder.")
    doc.save()
    print(f"DOCX saved to {args.output}")
